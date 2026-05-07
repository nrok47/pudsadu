#!/usr/bin/env python3
"""
สร้าง Excel รวม ครุภัณฑ์ + คลังสื่อ
พร้อมแก้วันที่ BE (1968/1969 → 2568/2569)
"""
import json
import openpyxl
from openpyxl import Workbook
from openpyxl.styles import (
    Font, PatternFill, Alignment, Border, Side, numbers
)
from openpyxl.utils import get_column_letter
from pathlib import Path
from datetime import datetime

# ─── Source files ───────────────────────────────────────────────────────────
EXCEL_SRC = Path(
    "/root/.claude/uploads/d67c8eb6-6ed9-4c01-812e-223d4b206ed0/"
    "bb0f7a59-___________________.xlsx"
)
EQUIP_JSON = Path("data/equipment.json")
MEDIA_JSON = Path("data/media.json")
OUT_FILE   = Path("data/รวมทะเบียน.xlsx")

# ─── Styles ─────────────────────────────────────────────────────────────────
BLUE   = "1F6FEB"
GREEN  = "1A7F37"
ORANGE = "E85D04"
PURPLE = "7048E8"
GRAY   = "6E7781"

def hdr_fill(hex_color: str) -> PatternFill:
    return PatternFill("solid", fgColor=hex_color)

def hdr_font(bold=True) -> Font:
    return Font(name="Angsana New", size=14, bold=bold, color="FFFFFF")

def body_font(bold=False) -> Font:
    return Font(name="Angsana New", size=13, bold=bold)

def thin_border() -> Border:
    s = Side(style="thin", color="D0D7DE")
    return Border(left=s, right=s, top=s, bottom=s)

def set_col_width(ws, col_widths: dict):
    for col, w in col_widths.items():
        ws.column_dimensions[col].width = w

def style_header_row(ws, row: int, fill_color: str, ncols: int):
    for col in range(1, ncols + 1):
        cell = ws.cell(row=row, column=col)
        cell.fill   = hdr_fill(fill_color)
        cell.font   = hdr_font()
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin_border()

def style_data_row(ws, row: int, ncols: int, even: bool):
    bg = "F6F8FA" if even else "FFFFFF"
    for col in range(1, ncols + 1):
        cell = ws.cell(row=row, column=col)
        cell.fill   = PatternFill("solid", fgColor=bg)
        cell.font   = body_font()
        cell.border = thin_border()
        cell.alignment = Alignment(vertical="center", wrap_text=False)

# ─── Date helpers ────────────────────────────────────────────────────────────
def fix_be_date(val) -> str:
    """แปลง datetime/string ปี 1960-1975 → Thai BE (+600) และ clean text
    Excel เก็บ '17/7/68' เป็น string '1968-07-17 00:00:00' หรือ datetime
    ต้องบวก 600 ปีเพื่อได้ปี พศ ที่ถูกต้อง (1968 → 2568)
    """
    import re as _re
    if val is None:
        return ""
    if isinstance(val, datetime):
        y = val.year
        be = y + 600 if 1960 <= y <= 1975 else y
        return f"{val.day}/{val.month}/{be}"
    s = str(val).strip().replace("\n", " ")
    # Handle "1968-07-17 00:00:00" or "1968-07-17" patterns
    m = _re.match(r"^(1\d{3})-(\d{2})-(\d{2})", s)
    if m:
        y, mo, d = int(m.group(1)), int(m.group(2)), int(m.group(3))
        if 1960 <= y <= 1975:
            return f"{d}/{mo}/{y + 600}"
    return s

# ─── Sheet 1: ภาพรวม (Summary) ──────────────────────────────────────────────
def build_summary(wb: Workbook, equip: list, media: list):
    ws = wb.create_sheet("📊 ภาพรวม")
    ws.sheet_view.showGridLines = False

    # Title
    ws.merge_cells("A1:H1")
    t = ws["A1"]
    t.value = "📊 ภาพรวมระบบทะเบียนฝ่าย"
    t.font  = Font(name="Angsana New", size=18, bold=True, color="1F6FEB")
    t.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 35

    # ─ ครุภัณฑ์ stats ─
    ws["A3"] = "🔧 สรุปครุภัณฑ์"
    ws["A3"].font = Font(name="Angsana New", size=14, bold=True, color=BLUE)
    ws.merge_cells("A3:H3")

    equip_headers = ["ประเภท", "จำนวน (รายการ)", "มูลค่ารวม (บาท)"]
    for ci, h in enumerate(equip_headers, 1):
        c = ws.cell(row=4, column=ci, value=h)
        c.fill = hdr_fill(BLUE); c.font = hdr_font()
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = thin_border()

    cat_map: dict[str, dict] = {}
    for e in equip:
        cat = e.get("category") or "ไม่ระบุ"
        if cat not in cat_map:
            cat_map[cat] = {"count": 0, "value": 0.0}
        cat_map[cat]["count"] += 1
        cat_map[cat]["value"] += e.get("unit_price") or 0.0

    r = 5
    for cat, info in sorted(cat_map.items(), key=lambda x: -x[1]["count"]):
        ws.cell(row=r, column=1, value=cat)
        ws.cell(row=r, column=2, value=info["count"])
        price_cell = ws.cell(row=r, column=3, value=info["value"])
        price_cell.number_format = '#,##0.00'
        style_data_row(ws, r, 3, r % 2 == 0)
        r += 1

    # total row
    total_eq   = len(equip)
    total_val  = sum(e.get("unit_price") or 0.0 for e in equip)
    ws.cell(row=r, column=1, value="รวมทั้งหมด").font = body_font(bold=True)
    ws.cell(row=r, column=2, value=total_eq).font = body_font(bold=True)
    v = ws.cell(row=r, column=3, value=total_val)
    v.number_format = '#,##0.00'; v.font = body_font(bold=True)
    style_header_row(ws, r, BLUE, 3)
    r += 2

    # ─ คลังสื่อ stats ─
    ws.merge_cells(f"A{r}:H{r}")
    ws[f"A{r}"] = "📚 สรุปคลังสื่อ"
    ws[f"A{r}"].font = Font(name="Angsana New", size=14, bold=True, color=GREEN)
    r += 1

    media_headers = ["กลุ่มงาน", "รายการ", "รับรวม (ชิ้น)", "จ่ายรวม (ชิ้น)", "คงเหลือ (ชิ้น)", "หมด", "ใกล้หมด"]
    for ci, h in enumerate(media_headers, 1):
        c = ws.cell(row=r, column=ci, value=h)
        c.fill = hdr_fill(GREEN); c.font = hdr_font()
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = thin_border()
    r += 1

    group_map: dict[str, dict] = {}
    for m in media:
        g = m.get("department") or "ไม่ระบุ"
        if g not in group_map:
            group_map[g] = {"count":0,"received":0,"issued":0,"balance":0,"out":0,"low":0}
        d = group_map[g]
        d["count"]    += 1
        d["received"] += m.get("received") or 0
        d["issued"]   += m.get("issued")   or 0
        d["balance"]  += m.get("balance")  or 0
        if m.get("status") == "หมด":     d["out"] += 1
        if m.get("status") == "ใกล้หมด": d["low"] += 1

    for dept, d in sorted(group_map.items()):
        row_data = [dept, d["count"], d["received"], d["issued"], d["balance"], d["out"], d["low"]]
        for ci, v in enumerate(row_data, 1):
            ws.cell(row=r, column=ci, value=v)
        style_data_row(ws, r, 7, r % 2 == 0)
        r += 1

    # totals
    tot = {"count":0,"received":0,"issued":0,"balance":0,"out":0,"low":0}
    for d in group_map.values():
        for k in tot: tot[k] += d[k]
    for ci, v in enumerate(["รวม", tot["count"], tot["received"], tot["issued"], tot["balance"], tot["out"], tot["low"]], 1):
        c = ws.cell(row=r, column=ci, value=v)
        c.font = body_font(bold=True)
    style_header_row(ws, r, GREEN, 7)

    set_col_width(ws, {"A":30,"B":14,"C":16,"D":16,"E":16,"F":8,"G":8})
    ws.row_dimensions[1].height = 35

# ─── Sheet 2: ครุภัณฑ์ ──────────────────────────────────────────────────────
def build_equip(wb: Workbook, equip: list):
    ws = wb.create_sheet("🔧 ครุภัณฑ์")
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "A3"

    ws.merge_cells("A1:O1")
    t = ws["A1"]
    t.value = f"🔧 ทะเบียนครุภัณฑ์  —  {len(equip)} รายการ"
    t.font  = Font(name="Angsana New", size=16, bold=True, color=BLUE)
    t.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 30

    headers = [
        "เลขครุภัณฑ์ FSN", "เลขครุภัณฑ์", "เลขครุภัณฑ์ GF",
        "ปีงบฯ", "ประเภท", "หมวด/Tags",
        "ชื่อ/คุณลักษณะ", "สถานที่", "ผู้รับผิดชอบ",
        "ยี่ห้อ", "รุ่น", "ราคา/หน่วย", "หน่วยนับ",
        "วันที่รับเข้า", "สถานะ",
    ]
    for ci, h in enumerate(headers, 1):
        c = ws.cell(row=2, column=ci, value=h)
    style_header_row(ws, 2, BLUE, len(headers))
    ws.row_dimensions[2].height = 22

    for ri, e in enumerate(equip, 3):
        row_data = [
            e.get("fsn"), e.get("equipment_id"), e.get("gf_id"),
            e.get("fiscal_year"), e.get("category"),
            "|".join(e.get("tags") or []),
            e.get("name"), e.get("location"), e.get("responsible"),
            e.get("brand"), e.get("model"),
            e.get("unit_price"), e.get("unit"),
            e.get("received_date"), e.get("status"),
        ]
        for ci, v in enumerate(row_data, 1):
            ws.cell(row=ri, column=ci, value=v)
        # price formatting
        price_cell = ws.cell(row=ri, column=12)
        if price_cell.value:
            price_cell.number_format = '#,##0.00'
        style_data_row(ws, ri, len(headers), ri % 2 == 0)

    col_widths = {"A":16,"B":16,"C":16,"D":8,"E":22,"F":20,
                  "G":35,"H":22,"I":25,"J":14,"K":14,
                  "L":14,"M":10,"N":14,"O":8}
    set_col_width(ws, col_widths)

# ─── Sheet 3: คลังสื่อ ──────────────────────────────────────────────────────
def build_media(wb: Workbook, media: list):
    ws = wb.create_sheet("📚 คลังสื่อ")
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "A3"

    ws.merge_cells("A1:J1")
    t = ws["A1"]
    t.value = f"📚 ทะเบียนคลังสื่อ  —  {len(media)} รายการ"
    t.font  = Font(name="Angsana New", size=16, bold=True, color=GREEN)
    t.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 30

    headers = ["รหัสสื่อ", "กลุ่มงาน", "กลุ่มงาน (ย่อ)", "ชื่อสื่อ / รายการ",
               "หน่วย", "รับรวม", "จ่ายรวม", "คงเหลือ", "สถานะ", "Link"]
    for ci, h in enumerate(headers, 1):
        c = ws.cell(row=2, column=ci, value=h)
    style_header_row(ws, 2, GREEN, len(headers))

    STATUS_COLOR = {"หมด": "FFEBE9", "ใกล้หมด": "FFF8C5", "มีเหลือ": "DAFBE1"}

    for ri, m in enumerate(media, 3):
        row_data = [
            m.get("media_id"), m.get("group_raw"), m.get("department"),
            m.get("name"), m.get("unit"),
            m.get("received"), m.get("issued"), m.get("balance"),
            m.get("status"), m.get("link"),
        ]
        for ci, v in enumerate(row_data, 1):
            ws.cell(row=ri, column=ci, value=v)
        style_data_row(ws, ri, len(headers), ri % 2 == 0)
        # colour status cell
        status_val = m.get("status") or ""
        bg = STATUS_COLOR.get(status_val, "FFFFFF")
        ws.cell(row=ri, column=9).fill = PatternFill("solid", fgColor=bg)

    set_col_width(ws, {"A":10,"B":22,"C":20,"D":55,"E":10,
                       "F":10,"G":10,"H":10,"I":10,"J":50})

# ─── Sheet 4: รับเข้า ────────────────────────────────────────────────────────
def build_stock_in(wb: Workbook, ws_src):
    ws = wb.create_sheet("📥 รับเข้า")
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "A3"

    ws.merge_cells("A1:I1")
    t = ws["A1"]
    t.value = "📥 รายการรับเข้า"
    t.font  = Font(name="Angsana New", size=16, bold=True, color=PURPLE)
    t.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 30

    headers = ["ลำดับ", "วันที่รับ (พศ)", "กลุ่มงาน", "รหัสสื่อ",
               "ชื่อสื่อ", "จำนวนรับ", "หน่วย", "สถานะ", "หมายเหตุ"]
    for ci, h in enumerate(headers, 1):
        ws.cell(row=2, column=ci, value=h)
    style_header_row(ws, 2, PURPLE, len(headers))

    ri = 3
    seq = 1
    media_lookup: dict[str, dict] = {}
    # pre-load media data
    try:
        with open("data/media.json", encoding="utf-8") as f:
            media = json.load(f)
        media_lookup = {m["media_id"]: m for m in media}
    except Exception:
        pass

    for row in ws_src.iter_rows(min_row=3, values_only=True):
        code = str(row[3]).strip() if row[3] else None
        if not code or not code.startswith("M"):
            continue
        qty  = row[5]
        name = media_lookup.get(code, {}).get("name", "")
        unit = media_lookup.get(code, {}).get("unit", "")
        row_data = [
            seq,
            fix_be_date(row[1]),
            str(row[2]).strip() if row[2] else "",
            code, name, qty, unit,
            str(row[7]).strip() if row[7] else "",
            str(row[8]).strip() if row[8] else "",
        ]
        for ci, v in enumerate(row_data, 1):
            ws.cell(row=ri, column=ci, value=v)
        style_data_row(ws, ri, len(headers), ri % 2 == 0)
        ri += 1; seq += 1

    set_col_width(ws, {"A":8,"B":18,"C":18,"D":10,"E":50,"F":10,"G":8,"H":10,"I":25})

# ─── Sheet 5: จ่ายออก ────────────────────────────────────────────────────────
def build_stock_out(wb: Workbook, ws_src):
    ws = wb.create_sheet("📤 จ่ายออก")
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "A3"

    ws.merge_cells("A1:J1")
    t = ws["A1"]
    t.value = "📤 รายการจ่ายออก (เบิก)"
    t.font  = Font(name="Angsana New", size=16, bold=True, color=ORANGE)
    t.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 30

    headers = ["ลำดับ", "วันที่จ่าย (พศ)", "เลขที่ใบเบิก", "กลุ่มงานเบิก",
               "รหัสสื่อ", "ชื่อสื่อ", "จำนวนจ่าย", "หน่วย", "สถานะ", "หมายเหตุ"]
    for ci, h in enumerate(headers, 1):
        ws.cell(row=2, column=ci, value=h)
    style_header_row(ws, 2, ORANGE, len(headers))

    media_lookup: dict[str, dict] = {}
    try:
        with open("data/media.json", encoding="utf-8") as f:
            media = json.load(f)
        media_lookup = {m["media_id"]: m for m in media}
    except Exception:
        pass

    ri = 3; seq = 1
    for row in ws_src.iter_rows(min_row=3, values_only=True):
        code = str(row[4]).strip() if row[4] else None
        if not code or not code.startswith("M"):
            continue
        qty  = row[6]
        name = media_lookup.get(code, {}).get("name", "")
        unit = media_lookup.get(code, {}).get("unit", "")
        row_data = [
            seq,
            fix_be_date(row[1]),
            str(row[2]).strip() if row[2] else "",
            str(row[3]).strip() if row[3] else "",
            code, name, qty, unit,
            str(row[8]).strip() if row[8] else "",
            str(row[9]).strip() if row[9] else "",
        ]
        for ci, v in enumerate(row_data, 1):
            ws.cell(row=ri, column=ci, value=v)
        style_data_row(ws, ri, len(headers), ri % 2 == 0)
        ri += 1; seq += 1

    set_col_width(ws, {"A":8,"B":18,"C":14,"D":18,"E":10,"F":50,"G":10,"H":8,"I":10,"J":25})

# ─── Main ────────────────────────────────────────────────────────────────────
def main():
    with open(EQUIP_JSON, encoding="utf-8") as f:
        equip = json.load(f)
    with open(MEDIA_JSON, encoding="utf-8") as f:
        media = json.load(f)

    src_wb = openpyxl.load_workbook(EXCEL_SRC, data_only=True)

    wb = Workbook()
    wb.remove(wb.active)   # remove default sheet

    build_summary(wb, equip, media)
    build_equip(wb, equip)
    build_media(wb, media)
    build_stock_in(wb, src_wb["📥 รับเข้า"])
    build_stock_out(wb, src_wb["📤 จ่ายออก"])

    wb.save(OUT_FILE)
    print(f"Saved → {OUT_FILE}")
    print(f"  ครุภัณฑ์ : {len(equip)} รายการ")
    print(f"  คลังสื่อ : {len(media)} รายการ")


if __name__ == "__main__":
    main()
